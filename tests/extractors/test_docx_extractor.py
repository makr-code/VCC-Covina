"""
Tests for DOCXExtractor

Creates a simple DOCX using python-docx and verifies extraction.
"""

import pytest
from pathlib import Path

from docx import Document

from ingestion.core.router import ChunkRouter
from ingestion.extractors.docx_extractor import DOCXExtractor


@pytest.mark.asyncio
async def test_docx_extractor_simple_text(tmp_path):
    # Create a simple DOCX file
    docx_path = tmp_path / "hello.docx"
    doc = Document()
    doc.add_paragraph("Hello DOCX World")
    doc.add_paragraph("Second line")
    doc.save(str(docx_path))

    # Setup router with DOCX extractor
    router = ChunkRouter()
    router.register(DOCXExtractor())

    # Extract
    result = await router.extract_file(docx_path, correlation_id="corr-docx", job_id="job-2")

    assert result.success is True
    assert len(result.chunks) == 1
    text = result.chunks[0].text
    assert "Hello DOCX World" in text
    assert "Second line" in text
    assert result.chunks[0].metadata.correlation_id == "corr-docx"
    assert result.chunks[0].metadata.job_id == "job-2"


@pytest.mark.asyncio
async def test_docx_extractor_missing_file(tmp_path):
    router = ChunkRouter()
    router.register(DOCXExtractor())

    missing = tmp_path / "missing.docx"
    res = await router.extract_file(missing)

    assert res.success is False
    assert "No extractor" not in res.error  # Format is supported
    assert "not found" in res.error.lower()
